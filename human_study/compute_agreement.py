#!/usr/bin/env python3
"""compute_agreement.py — Package D analysis for ArabAgentHallu (R2-6).

Usage:  python compute_agreement.py rater1.xlsx rater2.xlsx

Reads the two filled rater forms (sheet 'التقييم', data rows 3..62) and prints:
naturalness means±sd per language per rater, equivalence rate, Cohen's kappa
on equivalence, naturalness inter-rater correlation, and a ready-to-paste
summary sentence for the response letter.
"""
import sys, math
from openpyxl import load_workbook

def read(path):
    ws = load_workbook(path, data_only=True)['التقييم']
    rows = []
    for r in range(3, 63):
        nat_ar, nat_en, eq = (ws.cell(r, 5).value, ws.cell(r, 6).value,
                              ws.cell(r, 7).value)
        rows.append((ws.cell(r, 2).value, nat_ar, nat_en,
                     None if eq is None else str(eq).strip() in ('نعم', 'yes', 'Yes')))
    missing = [i + 1 for i, x in enumerate(rows) if None in x[1:3] or x[3] is None]
    if missing:
        print(f'⚠️ {path}: صفوف ناقصة: {missing}')
    return rows

def mean(xs): xs = list(xs); return sum(xs) / len(xs)
def sd(xs):
    xs = list(xs); m = mean(xs)
    return math.sqrt(sum((x - m) ** 2 for x in xs) / (len(xs) - 1))

def kappa(a, b):
    po = mean(1 if x == y else 0 for x, y in zip(a, b))
    pa, pb = mean(a), mean(b)
    pe = pa * pb + (1 - pa) * (1 - pb)
    return (po - pe) / (1 - pe) if pe < 1 else float('nan')

def pearson(a, b):
    ma, mb = mean(a), mean(b)
    num = sum((x - ma) * (y - mb) for x, y in zip(a, b))
    den = math.sqrt(sum((x - ma) ** 2 for x in a) * sum((y - mb) ** 2 for y in b))
    return num / den if den else float('nan')

r1, r2 = read(sys.argv[1]), read(sys.argv[2])
assert [x[0] for x in r1] == [x[0] for x in r2], 'ترتيب العناصر مختلف بين الملفين'

for name, rr in (('Rater 1', r1), ('Rater 2', r2)):
    a = [x[1] for x in rr]; e = [x[2] for x in rr]; q = [x[3] for x in rr]
    print(f'{name}: AR naturalness {mean(a):.2f}±{sd(a):.2f} | '
          f'EN {mean(e):.2f}±{sd(e):.2f} | equivalent {100*mean(q):.1f}%')

q1, q2 = [x[3] for x in r1], [x[3] for x in r2]
k = kappa(q1, q2)
both = mean(1 if a and b else 0 for a, b in zip(q1, q2))
par = pearson([x[1] for x in r1], [x[1] for x in r2])
pen = pearson([x[2] for x in r1], [x[2] for x in r2])
print(f'Equivalence: both-yes {100*both:.1f}% | Cohen kappa {k:.3f}')
print(f'Naturalness inter-rater r: AR {par:.3f} | EN {pen:.3f}')
print()
print('--- sentence for the response letter / paper ---')
print(f'Two Arabic-English bilingual raters independently scored 60 randomly '
      f'sampled item pairs (30 hallucinated / 30 clean, stratified by category, '
      f'seed 13). Mean naturalness was '
      f'{mean([x[1] for x in r1+r2]):.2f}/5 (Arabic) and '
      f'{mean([x[2] for x in r1+r2]):.2f}/5 (English); '
      f'{100*both:.1f}% of pairs were judged semantically equivalent by both '
      f'raters (Cohen kappa = {k:.2f}).')
